import streamlit as st
import pdfplumber
import pandas as pd
import datetime
import io
import re
import random
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(page_title="Gerador de Escala (Final)", page_icon="🏥", layout="centered")

# --- CORES E CONSTANTES ---
COR_AZUL_CLARO = "CFE2F3"
COR_ROSA_CLARO = "F4CCCC"
COR_ROSA_ESCURO = "EA9999"

# --- FUNÇÕES DE LÓGICA DE EXTRAÇÃO ---

def formatar_nome(nome_completo):
    if not isinstance(nome_completo, str): return ""
    partes = nome_completo.split()
    ignorar = ["ENF", "ENFERMEIRO", "CONTRATO", "EFETIVO", "TEC", "TECNICO", "MÉDIO", "MEDIO", "VÍNCULO", "FUNÇÃO", "COREN", "COREN-AP"]
    partes = [p for p in partes if p.upper() not in ignorar and len(p) > 2]
    if len(partes) > 1: return f"{partes[0]} {partes[-1]}".title()
    elif len(partes) == 1: return partes[0].title()
    return ""

def limpar_valor(val):
    return str(val).strip() if val is not None else ""

def detectar_metadados(pdf, nome_arquivo):
    nome_upper = nome_arquivo.upper()
    tipo = "ENFERMEIROS"
    if "TEC" in nome_upper or "TÉC" in nome_upper: tipo = "TÉCNICOS"
    elif "ENF" in nome_upper: tipo = "ENFERMEIROS"
    
    texto_completo = ""
    try: texto_completo = pdf.pages[0].extract_text().upper()
    except: pass
    
    meses = {r'JANEIRO':1, r'FEVEREIRO':2, r'MARÇO':3, r'MARCO':3, r'ABRIL':4, r'MAIO':5, r'JUNHO':6, r'JULHO':7, r'AGOSTO':8, r'SETEMBRO':9, r'OUTUBRO':10, r'NOVEMBRO':11, r'DEZEMBRO':12}
    mes_detectado = 1 # Default
    for r_mes, n_mes in meses.items():
        if r_mes in texto_completo:
            mes_detectado = n_mes
            break
            
    ano_detectado = 2026
    match_ano = re.search(r'202[4-9]', texto_completo)
    if match_ano: ano_detectado = int(match_ano.group(0))

    return tipo, ano_detectado, mes_detectado

def processar_pdf(file_obj, nome_arquivo):
    dados = []
    with pdfplumber.open(file_obj) as pdf:
        tipo, ano, mes = detectar_metadados(pdf, nome_arquivo)
        
        for page in pdf.pages:
            tabelas = page.extract_tables()
            for tabela in tabelas:
                df = pd.DataFrame(tabela)
                
                idx_cabecalho = -1
                mapa_dias = {}
                
                for idx, row in df.iterrows():
                    numeros_validos = []
                    for c, val in enumerate(row):
                        v_str = limpar_valor(val)
                        if v_str.isdigit() and 1 <= int(v_str) <= 31:
                            numeros_validos.append(int(v_str))
                    
                    if len(numeros_validos) >= 5:
                        idx_cabecalho = idx
                        
                        # Trava de Virada de Mês
                        ultimo_dia_visto = 0
                        for c, v in enumerate(row):
                            vl = limpar_valor(v)
                            if vl.isdigit():
                                dia_num = int(vl)
                                if 1 <= dia_num <= 31:
                                    if dia_num < ultimo_dia_visto and ultimo_dia_visto > 20:
                                        continue 
                                    mapa_dias[c] = dia_num
                                    ultimo_dia_visto = dia_num
                        break
                
                if idx_cabecalho == -1: continue
                
                df_dados = df.iloc[idx_cabecalho+1:].copy()
                col_nome = 1
                for i, col in enumerate(df.columns):
                    if "NOME" in str(col).upper(): col_nome = i; break
                
                for _, row in df_dados.iterrows():
                    if col_nome >= len(row): continue
                    nome = formatar_nome(limpar_valor(row.iloc[col_nome]))
                    if len(nome) < 3 or "TURNO" in nome.upper(): continue
                    
                    for c_idx, dia in mapa_dias.items():
                        if c_idx < len(row):
                            turno = limpar_valor(row.iloc[c_idx]).upper()
                            validos = ["M", "T", "N", "N1", "N2", "D", "SD", "SN", "LP", "MT", "TM", "MD"]
                            eh_valido = False
                            if len(turno) < 7:
                                for v in validos:
                                    if v in turno: 
                                        eh_valido = True
                                        break
                            
                            if eh_valido:
                                dados.append({"DIA": dia, "TURNO": turno, "NOME": nome})
                                
    if not dados:
        return pd.DataFrame(columns=['DIA', 'TURNO', 'NOME']), tipo, ano, mes
    return pd.DataFrame(dados), tipo, ano, mes

# --- FUNÇÕES DE WORD / LAYOUT ---

def definir_cor_fundo(celula, cor_hex):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), cor_hex))
    celula._tc.get_or_add_tcPr().append(shading_elm)

def formatar_texto(run, tamanho=10, negrito=False, alinhamento=None):
    font = run.font
    font.size = Pt(tamanho)
    font.bold = negrito
    font.name = 'Arial'

def set_col_widths(table):
    # Larguras Aproximadas (Total A4 ~16-17cm úteis)
    widths = [Cm(2.0), Cm(0.8), Cm(6.5), Cm(3.5), Cm(3.5)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = width

def definir_funcoes_aleatorias(df_turno):
    """Retorna uma lista de funções (Volante/Classificador) alinhada ao dataframe"""
    if df_turno.empty: return []
    
    total = len(df_turno)
    funcoes = ["Classificador"] * total
    
    # Escolhe um aleatório para ser Volante
    if total > 0:
        idx_volante = random.randint(0, total - 1)
        funcoes[idx_volante] = "Volante"
        
    return funcoes

def adicionar_bloco_turno_enf(table, df_filtrado, nome_turno, cor_lateral, start_row_idx):
    qtd = max(1, len(df_filtrado))
    
    # Gera funções se houver enfermeiros
    funcoes = definir_funcoes_aleatorias(df_filtrado)
    
    for _ in range(qtd): table.add_row()
        
    # Mescla Lateral (Turno Label)
    c1 = table.rows[start_row_idx].cells[0]
    c2 = table.rows[start_row_idx + qtd - 1].cells[0]
    merged = c1.merge(c2)
    merged.text = nome_turno
    merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    formatar_texto(merged.paragraphs[0].runs[0], negrito=True, tamanho=9)
    definir_cor_fundo(merged, cor_lateral)
    
    if df_filtrado.empty:
        for i in range(1, 5): table.rows[start_row_idx].cells[i].text = "-"
    else:
        for i, (_, row) in enumerate(df_filtrado.iterrows()):
            r = table.rows[start_row_idx + i]
            
            # Turno
            r.cells[1].text = row['TURNO']
            r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            formatar_texto(r.cells[1].paragraphs[0].runs[0], negrito=True, tamanho=9)
            
            # Nome
            r.cells[2].text = row['NOME']
            formatar_texto(r.cells[2].paragraphs[0].runs[0], tamanho=9)
            
            # Função
            r.cells[3].text = funcoes[i]
            r.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            formatar_texto(r.cells[3].paragraphs[0].runs[0], tamanho=8)

            # Trocas (Vazio)
            r.cells[4].text = ""

    return start_row_idx + qtd

def adicionar_bloco_turno_tec(table, df_filtrado, nome_turno, cor_lateral, start_row_idx):
    qtd = max(1, len(df_filtrado))
    for _ in range(qtd): table.add_row()
        
    c1 = table.rows[start_row_idx].cells[0]
    c2 = table.rows[start_row_idx + qtd - 1].cells[0]
    merged = c1.merge(c2)
    merged.text = nome_turno
    merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    formatar_texto(merged.paragraphs[0].runs[0], negrito=True, tamanho=9)
    definir_cor_fundo(merged, cor_lateral)
    
    if df_filtrado.empty:
        table.rows[start_row_idx].cells[1].text = "-"
        table.rows[start_row_idx].cells[2].text = "-"
    else:
        for i, (_, row) in enumerate(df_filtrado.iterrows()):
            r = table.rows[start_row_idx + i]
            
            # Turno
            r.cells[1].text = row['TURNO']
            r.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            formatar_texto(r.cells[1].paragraphs[0].runs[0], negrito=True, tamanho=9)
            
            # Nome
            r.cells[2].text = row['NOME']
            formatar_texto(r.cells[2].paragraphs[0].runs[0], tamanho=9)
            
            # Trocas (Mescla Col 3 e 4 para ficar maior)
            c_troca = r.cells[3].merge(r.cells[4])
            c_troca.text = ""
            
    return start_row_idx + qtd

def gerar_docx_completo(df_enf, df_tec, ano, mes):
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    
    # Tenta ajustar margens para caber mais coisa
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    df_enf['DIA'] = pd.to_numeric(df_enf['DIA'], errors='coerce')
    df_tec['DIA'] = pd.to_numeric(df_tec['DIA'], errors='coerce')
    
    dias_enf = set(df_enf['DIA'].dropna().unique())
    dias_tec = set(df_tec['DIA'].dropna().unique())
    dias_totais = sorted(list(dias_enf | dias_tec))
    
    dias_semana = {0: 'SEGUNDA-FEIRA', 1: 'TERÇA-FEIRA', 2: 'QUARTA-FEIRA', 3: 'QUINTA-FEIRA', 4: 'SEXTA-FEIRA', 5: 'SÁBADO', 6: 'DOMINGO'}

    for dia in dias_totais:
        try:
            dt = datetime.date(ano, mes, int(dia))
            txt_data = f"{dt.strftime('%d/%m/%Y')} {dias_semana[dt.weekday()]}"
        except: continue
        
        # Cria tabela com 5 colunas
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        table.autofit = False 
        set_col_widths(table)
        
        # Data Header
        r = table.rows[0]
        c = r.cells[0].merge(r.cells[4])
        c.text = txt_data
        definir_cor_fundo(c, COR_AZUL_CLARO)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        formatar_texto(c.paragraphs[0].runs[0], tamanho=11, negrito=True)
        
        # === ENFERMEIROS ===
        r = table.add_row()
        c = r.cells[0].merge(r.cells[4])
        c.text = "ENFERMEIROS"
        definir_cor_fundo(c, COR_ROSA_CLARO)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        formatar_texto(c.paragraphs[0].runs[0], tamanho=10, negrito=True)
        
        # Cabeçalho Colunas Enfermeiro
        r_head = table.add_row()
        col_names = ["", "Turno", "Nome", "Função", "Trocas"]
        for idx, nome in enumerate(col_names):
            r_head.cells[idx].text = nome
            r_head.cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx > 0: formatar_texto(r_head.cells[idx].paragraphs[0].runs[0], tamanho=8, negrito=True)
        
        sub_enf = df_enf[df_enf['DIA'] == dia].sort_values('TURNO')
        not_enf = sub_enf[sub_enf['TURNO'].str.contains('N')]
        diu_enf = sub_enf[~sub_enf.index.isin(not_enf.index)]
        
        idx = 3
        idx = adicionar_bloco_turno_enf(table, diu_enf, "DIURNO", COR_AZUL_CLARO, idx)
        idx = adicionar_bloco_turno_enf(table, not_enf, "NOTURNO", COR_ROSA_ESCURO, idx)

        # === TÉCNICOS ===
        r = table.add_row()
        c = r.cells[0].merge(r.cells[4])
        c.text = "TÉCNICOS"
        definir_cor_fundo(c, COR_ROSA_CLARO)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        formatar_texto(c.paragraphs[0].runs[0], tamanho=10, negrito=True)
        
        # Cabeçalho Colunas Técnicos
        r_head = table.add_row()
        col_names_tec = ["", "Turno", "Nome", "Trocas", ""]
        for idx, nome in enumerate(col_names_tec):
            if idx == 4: continue # Pula ultima
            cell = r_head.cells[idx]
            if idx == 3: cell = cell.merge(r_head.cells[4])
            
            cell.text = nome
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx > 0: formatar_texto(cell.paragraphs[0].runs[0], tamanho=8, negrito=True)
            
        idx += 2
        
        sub_tec = df_tec[df_tec['DIA'] == dia].sort_values('TURNO')
        not_tec = sub_tec[sub_tec['TURNO'].str.contains('N')]
        diu_tec = sub_tec[~sub_tec.index.isin(not_tec.index)]
        
        idx = adicionar_bloco_turno_tec(table, diu_tec, "DIURNO", COR_AZUL_CLARO, idx)
        idx = adicionar_bloco_turno_tec(table, not_tec, "NOTURNO", COR_ROSA_ESCURO, idx)
        
        doc.add_paragraph("")

    return doc

# --- INTERFACE ---
st.title("🏥 Gerador de Escala (Layout Ajustado)")
st.markdown("""
- **Enfermeiros:** Define automaticamente 1 Volante (aleatório) e o resto Classificador. Coluna de Trocas inclusa.
- **Técnicos:** Coluna de Trocas expandida.
""")

uploaded_files = st.file_uploader("Arraste os arquivos aqui", type=["pdf"], accept_multiple_files=True)

if uploaded_files:
    if st.button("🚀 Processar"):
        dfs = {'ENFERMEIROS': pd.DataFrame(), 'TÉCNICOS': pd.DataFrame()}
        meta_ano, meta_mes = 2026, 1
        
        for f in uploaded_files:
            df, tipo, ano, mes = processar_pdf(f, f.name)
            if not df.empty:
                dfs[tipo] = df
                meta_ano, meta_mes = ano, mes
                st.success(f"Lido: {tipo} ({mes}/{ano})")

        doc = gerar_docx_completo(dfs['ENFERMEIROS'], dfs['TÉCNICOS'], meta_ano, meta_mes)
        
        bio = io.BytesIO()
        doc.save(bio)
        
        st.download_button(
            label="📥 Baixar DOCX",
            data=bio.getvalue(),
            file_name=f"Escala_Final_{meta_mes}_{meta_ano}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
